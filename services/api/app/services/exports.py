from __future__ import annotations

import csv
import io
import json
import zipfile
from pathlib import Path
from uuid import UUID

from openpyxl import Workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.errors import AppError, ErrorCodes
from app.storage.service import get_storage
from packages.db.enums import AssetType, ExportStatus
from packages.db.models import Export, MediaAsset, Observation, PlateDetection, ProcessingJob, VehicleTrack, Video


def _rows(db: Session, job: ProcessingJob) -> list[tuple[PlateDetection, VehicleTrack, Video]]:
    video = db.get(Video, job.video_id)
    stmt = (
        select(PlateDetection, VehicleTrack)
        .join(VehicleTrack)
        .where(VehicleTrack.processing_job_id == job.id)
        .order_by(PlateDetection.first_seen_seconds)
    )
    return [(d, t, video) for d, t in db.execute(stmt).all()]  # type: ignore[misc]


def create_export(db: Session, job_id: UUID, export_type: str) -> Export:
    job = db.get(ProcessingJob, job_id)
    if not job:
        raise AppError(ErrorCodes.NOT_FOUND, "Job not found", status_code=404)
    export = Export(processing_job_id=job.id, export_type=export_type, status=ExportStatus.PENDING.value)
    db.add(export)
    db.flush()
    storage = get_storage()
    try:
        if export_type == "csv":
            key, mime, size = _csv(db, job, storage)
            atype = AssetType.EXPORT_CSV.value
        elif export_type == "xlsx":
            key, mime, size = _xlsx(db, job, storage)
            atype = AssetType.EXPORT_XLSX.value
        elif export_type == "json":
            key, mime, size = _json(db, job, storage)
            atype = AssetType.EXPORT_JSON.value
        elif export_type == "evidence-zip":
            key, mime, size = _zip(db, job, storage)
            atype = AssetType.EXPORT_ZIP.value
        else:
            raise AppError(ErrorCodes.VALIDATION_ERROR, f"Unknown export type {export_type}")
        asset = MediaAsset(
            processing_job_id=job.id,
            asset_type=atype,
            storage_key=key,
            mime_type=mime,
            file_size=size,
        )
        db.add(asset)
        db.flush()
        export.storage_key = key
        export.media_asset_id = asset.id
        export.status = ExportStatus.READY.value
    except AppError:
        export.status = ExportStatus.FAILED.value
        export.error_message = "Export failed"
        raise
    except Exception as exc:
        export.status = ExportStatus.FAILED.value
        export.error_message = str(exc)
        raise AppError(ErrorCodes.EXPORT_FAILED, "Export failed", status_code=500) from exc
    db.commit()
    db.refresh(export)
    return export


def _csv(db: Session, job: ProcessingJob, storage) -> tuple[str, str, int]:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(
        [
            "plate_number",
            "raw_ocr_text",
            "vehicle_type",
            "track_id",
            "ocr_confidence",
            "plate_detection_confidence",
            "first_seen_seconds",
            "last_seen_seconds",
            "best_frame_number",
            "status",
            "source_video",
            "processing_job_id",
        ]
    )
    for det, track, video in _rows(db, job):
        writer.writerow(
            [
                det.normalized_plate_text or "",
                det.raw_ocr_text or "",
                track.vehicle_type,
                track.track_id,
                det.ocr_confidence,
                det.plate_detection_confidence,
                det.first_seen_seconds,
                det.last_seen_seconds,
                det.best_frame_number,
                det.status,
                video.original_filename if video else "",
                str(job.id),
            ]
        )
    data = buf.getvalue().encode("utf-8")
    key = storage.save("exports", f"job-{job.id}.csv", data)
    return key, "text/csv", len(data)


def _xlsx(db: Session, job: ProcessingJob, storage) -> tuple[str, str, int]:
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    summary.append(["job_id", str(job.id)])
    summary.append(["status", job.status])
    summary.append(["profile", job.processing_profile])
    summary.append(["device", job.actual_device])
    summary.append(["vehicles", job.vehicles_detected])
    summary.append(["plates", job.plates_detected])

    unique = wb.create_sheet("Unique Detections")
    unique.append(
        [
            "plate_number",
            "raw_ocr_text",
            "vehicle_type",
            "track_id",
            "ocr_confidence",
            "status",
            "first_seen_seconds",
            "last_seen_seconds",
        ]
    )
    for det, track, _video in _rows(db, job):
        unique.append(
            [
                det.normalized_plate_text,
                det.raw_ocr_text,
                track.vehicle_type,
                track.track_id,
                det.ocr_confidence,
                det.status,
                det.first_seen_seconds,
                det.last_seen_seconds,
            ]
        )

    raw = wb.create_sheet("Raw Observations")
    raw.append(["frame", "timestamp", "raw", "normalized", "ocr_confidence", "track_id"])
    obs = db.execute(
        select(Observation, PlateDetection, VehicleTrack)
        .join(PlateDetection, Observation.plate_detection_id == PlateDetection.id)
        .join(VehicleTrack, PlateDetection.vehicle_track_id == VehicleTrack.id)
        .where(VehicleTrack.processing_job_id == job.id)
        .order_by(Observation.frame_number)
    ).all()
    for observation, det, track in obs:
        raw.append(
            [
                observation.frame_number,
                observation.timestamp_seconds,
                observation.raw_ocr_text,
                observation.normalized_plate_text,
                observation.ocr_confidence,
                track.track_id,
            ]
        )

    hist = wb.create_sheet("Confidence History")
    hist.append(["detection_id", "frame", "ocr_confidence", "normalized"])
    for observation, det, _track in obs:
        hist.append([str(det.id), observation.frame_number, observation.ocr_confidence, observation.normalized_plate_text])

    buf = io.BytesIO()
    wb.save(buf)
    data = buf.getvalue()
    key = storage.save("exports", f"job-{job.id}.xlsx", data)
    return key, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", len(data)


def _json(db: Session, job: ProcessingJob, storage) -> tuple[str, str, int]:
    payload = {
        "job_id": str(job.id),
        "status": job.status,
        "metrics": job.metrics,
        "detections": [],
    }
    for det, track, video in _rows(db, job):
        payload["detections"].append(
            {
                "id": str(det.id),
                "plate": det.normalized_plate_text,
                "raw": det.raw_ocr_text,
                "track_id": track.track_id,
                "vehicle_type": track.vehicle_type,
                "ocr_confidence": det.ocr_confidence,
                "status": det.status,
                "source_video": video.original_filename if video else None,
            }
        )
    data = json.dumps(payload, indent=2).encode("utf-8")
    key = storage.save("exports", f"job-{job.id}.json", data)
    return key, "application/json", len(data)


def _zip(db: Session, job: ProcessingJob, storage) -> tuple[str, str, int]:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        csv_key, _, _ = _csv(db, job, storage)
        csv_path = storage.get_path(csv_key)
        zf.write(csv_path, "metadata.csv")
        summary = {
            "job_id": str(job.id),
            "status": job.status,
            "error_code": job.error_code,
            "error_message": job.error_message,
            "metrics": job.metrics,
        }
        zf.writestr("processing_summary.json", json.dumps(summary, indent=2))
        for det, track, _video in _rows(db, job):
            folder = f"track_{track.track_id}_{det.normalized_plate_text or 'unknown'}"
            for asset_id, name in (
                (det.full_frame_asset_id, "full_frame.jpg"),
                (det.vehicle_crop_asset_id, "vehicle.jpg"),
                (det.plate_crop_asset_id, "plate.jpg"),
            ):
                if not asset_id:
                    continue
                asset = db.get(MediaAsset, asset_id)
                if not asset:
                    continue
                try:
                    path = storage.get_path(asset.storage_key)
                    zf.write(path, f"{folder}/{name}")
                except Exception:
                    continue
    data = buf.getvalue()
    key = storage.save("exports", f"job-{job.id}.zip", data)
    return key, "application/zip", len(data)
